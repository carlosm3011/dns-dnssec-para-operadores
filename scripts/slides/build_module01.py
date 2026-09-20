#!/usr/bin/env python3
"""Build 01-Introduccion.pptx from the lacnic46.pptx template."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
import copy

RED = RGBColor(0xC5, 0x34, 0x25)
DARK = RGBColor(0x33, 0x33, 0x33)
GRAY = RGBColor(0x66, 0x66, 0x66)
LIGHTGRAY = RGBColor(0xF3, 0xF3, 0xF3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"

prs = Presentation("lacnic46.pptx")

layouts = {l.name: l for l in prs.slide_masters[0].slide_layouts}


def get_placeholder(slide, ph_type=None, idx=None):
    for ph in slide.placeholders:
        if idx is not None and ph.placeholder_format.idx == idx:
            return ph
        if ph_type is not None and str(ph.placeholder_format.type) == ph_type:
            return ph
    return None


def set_title(slide, text):
    slide.shapes.title.text_frame.text = text


def set_body_bullets(placeholder, items, size=16, space_after=10, bullet_first=True):
    """items: list of str, or (str, level) tuples.

    bullet_first: when True (flat bullet lists), strip any inherited
    <a:buNone/> on paragraph 0 so the first line bullets like the rest.
    Set False when item 0 is meant to read as an un-bulleted header
    (e.g. two-column layouts where the first line is a column title).
    """
    tf = placeholder.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if i == 0 and bullet_first:
            # Some pre-existing template placeholders carry an explicit
            # <a:pPr marL="0" indent="0" ...><a:buNone/></a:pPr> baked onto
            # their first (only) paragraph, which both kills the bullet and
            # zeroes the indent that gives it its hanging space. Drop the
            # whole override so this paragraph is built fresh, matching the
            # ones we add below via p.level.
            pPr = p._pPr
            if pPr is not None:
                p._p.remove(pPr)
        p.level = level
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size - (2 if level > 0 else 0))
        run.font.name = FONT
        run.font.color.rgb = DARK
        p.space_after = Pt(space_after)


def add_textbox(slide, left, top, width, height, text, size=14, bold=False,
                 color=DARK, align=PP_ALIGN.LEFT, font=FONT, italic=False):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font
    run.font.color.rgb = color
    return tb


def add_multiline_textbox(slide, left, top, width, height, lines, size=13,
                           color=DARK, font=FONT, line_spacing=1.15, align=PP_ALIGN.LEFT,
                           italic=False, bold=False):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.name = font
        run.font.color.rgb = color
        run.font.italic = italic
        run.font.bold = bold
        p.line_spacing = line_spacing
    return tb


def add_box(slide, left, top, width, height, text, fill=RED, text_color=WHITE,
            size=13, bold=True, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE):
    shp = slide.shapes.add_shape(shape_type, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(45700)
    tf.margin_right = Emu(45700)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = FONT
    run.font.color.rgb = text_color
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return shp


def add_arrow(slide, x1, y1, x2, y2, color=GRAY, width=Pt(2)):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    conn.line.color.rgb = color
    conn.line.width = width
    conn.line.end_arrowhead = "block" if hasattr(conn.line, "end_arrowhead") else None
    # python-pptx has no direct arrowhead API; patch XML
    ln = conn.line._get_or_add_ln()
    tail = ln.makeelement(qn('a:tailEnd'), {'type': 'triangle'})
    ln.append(tail)
    return conn


SLIDE_W = prs.slide_width
SLIDE_H = prs.slide_height

# ---------------------------------------------------------------------------
# Slide 1 (existing) -- Title
# ---------------------------------------------------------------------------
s1 = prs.slides[0]
title_ph = get_placeholder(s1, "CENTER_TITLE") or s1.shapes.title
title_ph.text_frame.paragraphs[0].runs[0].text = "Módulo 1: Introducción a DNSSEC"
sub_ph = get_placeholder(s1, idx=1)
sub_ph.text_frame.paragraphs[0].text = ""
p = sub_ph.text_frame.paragraphs[0]
run = p.add_run()
run.text = "Curso de DNS y DNSSEC para Operadores  —  Carlos Martinez"
run.font.size = Pt(18)
run.font.name = FONT

# ---------------------------------------------------------------------------
# Slide 2 (existing) -- Agenda
# ---------------------------------------------------------------------------
s2 = prs.slides[1]
set_title(s2, "Agenda del módulo")
body2 = get_placeholder(s2, idx=1)
set_body_bullets(body2, [
    "Repaso breve de DNS: zonas, RRsets y delegación",
    "Fundamentos de criptografía: clave pública, hashes, firma digital",
    "¿Qué es DNSSEC?",
    "Qué problemas resuelve — y qué NO resuelve",
    "Cadena de confianza",
    "Conceptos clave y resumen",
], size=18, space_after=14)

print("Base slides 1-2 done. Slide count:", len(prs.slides))

# ---------------------------------------------------------------------------
# Slide 3 -- TITLE_ONLY: DNS en breve
# ---------------------------------------------------------------------------
s3 = prs.slides.add_slide(layouts["TITLE_ONLY"])
set_title(s3, "DNS en breve: zonas y RRsets")
add_multiline_textbox(s3, Inches(0.4), Inches(1.15), Inches(9.2), Inches(1.0), [
    "DNS traduce nombres (example.com) a datos como direcciones IP. Los datos se",
    "organizan en zonas; cada zona agrupa RRsets — conjuntos de registros del",
    "mismo nombre y tipo.",
], size=15)

code_bg = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.4), Inches(2.25), Inches(9.2), Inches(2.9))
code_bg.fill.solid()
code_bg.fill.fore_color.rgb = LIGHTGRAY
code_bg.line.fill.background()
code_bg.shadow.inherit = False
code_bg.text_frame.word_wrap = True
code_lines = [
    "example.com.        3600  IN  SOA   ns1.example.com. admin.example.com. (",
    "                                     2026071001 ; serial",
    "                                     3600       ; refresh )",
    "example.com.        3600  IN  NS    ns1.example.com.",
    "example.com.        3600  IN  NS    ns2.example.com.",
    "example.com.        3600  IN  A     93.184.216.34",
    "www.example.com.    3600  IN  CNAME example.com.",
]
tf = code_bg.text_frame
tf.margin_left = Inches(0.25)
tf.margin_top = Inches(0.2)
tf.vertical_anchor = MSO_ANCHOR.TOP
for i, line in enumerate(code_lines):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = line
    run.font.size = Pt(12)
    run.font.name = "Courier New"
    run.font.color.rgb = DARK

# ---------------------------------------------------------------------------
# Slide 4 -- TITLE_ONLY: Delegación de zonas (diagram)
# ---------------------------------------------------------------------------
s4 = prs.slides.add_slide(layouts["TITLE_ONLY"])
set_title(s4, "Las zonas se delegan")
by = Inches(2.3)
bh = Inches(0.9)
bw = Inches(2.0)
gap = Inches(0.4)
x0 = Inches(0.4)
add_box(s4, x0, by, bw, bh, "Raíz (.)", fill=RED)
add_box(s4, x0 + bw + gap, by, bw, bh, "TLD (.com)", fill=RED)
add_box(s4, x0 + 2 * (bw + gap), by, bw, bh, "example.com", fill=RED)
add_box(s4, x0 + 3 * (bw + gap), by, bw, bh, "A 93.184.216.34", fill=DARK, size=12)
for i in range(3):
    x1 = x0 + i * (bw + gap) + bw
    x2 = x1 + gap
    add_arrow(s4, x1, by + bh // 2, x2, by + bh // 2)
add_multiline_textbox(s4, Inches(0.5), Inches(3.7), Inches(9.0), Inches(0.5),
                       ["Un resolver sigue la cadena de delegaciones, servidor por servidor, hasta obtener la respuesta."],
                       size=14, align=PP_ALIGN.LEFT)
add_multiline_textbox(s4, Inches(0.5), Inches(4.3), Inches(9.0), Inches(1.0), [
    "El problema: nada en este mecanismo prueba que una respuesta sea auténtica.",
], size=15, color=RED)
s4.shapes[-1].text_frame.paragraphs[0].runs[0].font.bold = True

# ---------------------------------------------------------------------------
# Slide 5 -- MAIN_POINT: statement
# ---------------------------------------------------------------------------
s5 = prs.slides.add_slide(layouts["MAIN_POINT"])
set_title(s5, "Un resolver acepta lo que recibe.\nNada prueba que una respuesta sea auténtica.")

# ---------------------------------------------------------------------------
# Slide 6 -- SECTION_TITLE_AND_DESCRIPTION: Fundamentos de criptografía
# ---------------------------------------------------------------------------
s6 = prs.slides.add_slide(layouts["SECTION_TITLE_AND_DESCRIPTION"])
set_title(s6, "Fundamentos de criptografía")
sub6 = get_placeholder(s6, idx=1)
if sub6:
    sub6.text_frame.paragraphs[0].text = ""
    r = sub6.text_frame.paragraphs[0].add_run()
    r.text = "Tres conceptos que DNSSEC usa todo el tiempo"
    r.font.size = Pt(18)
    r.font.name = FONT
body6 = get_placeholder(s6, idx=2)
if body6:
    set_body_bullets(body6, [
        "Clave pública / clave privada",
        "Hashes (digest)",
        "Firma digital",
    ], size=16)

# ---------------------------------------------------------------------------
# Slide 7 -- TITLE_AND_TWO_COLUMNS: clave pública vs privada
# ---------------------------------------------------------------------------
s7 = prs.slides.add_slide(layouts["TITLE_AND_TWO_COLUMNS"])
set_title(s7, "Criptografía de clave pública")
col1 = get_placeholder(s7, idx=1)
col2 = get_placeholder(s7, idx=2)
set_body_bullets(col1, [
    "Clave privada",
    ("Secreta — nunca se comparte", 1),
    ("Se usa para FIRMAR", 1),
], size=17, space_after=12)
col1.text_frame.paragraphs[0].runs[0].font.bold = True
col1.text_frame.paragraphs[0].runs[0].font.color.rgb = RED
set_body_bullets(col2, [
    "Clave pública",
    ("Se distribuye libremente", 1),
    ("Se usa para VERIFICAR", 1),
    ("Publicada en la zona (DNSKEY)", 1),
], size=17, space_after=12)
col2.text_frame.paragraphs[0].runs[0].font.bold = True
col2.text_frame.paragraphs[0].runs[0].font.color.rgb = RED

# ---------------------------------------------------------------------------
# Slide 8 -- TITLE_AND_BODY: Hashes, tres propiedades
# ---------------------------------------------------------------------------
s8 = prs.slides.add_slide(layouts["TITLE_AND_BODY"])
set_title(s8, "Hashes: tres propiedades")
body8 = get_placeholder(s8, idx=1)
set_body_bullets(body8, [
    "Determinismo",
    ("El mismo dato siempre produce el mismo hash", 1),
    "Efecto avalancha",
    ("Cambiar un solo bit cambia por completo el hash resultante", 1),
    "Resistencia a colisiones y a preimagen",
    ("Inviable encontrar dos datos con el mismo hash, o reconstruir el original", 1),
], size=17, space_after=10)
for i in (0, 2, 4):
    body8.text_frame.paragraphs[i].runs[0].font.bold = True
    body8.text_frame.paragraphs[i].runs[0].font.color.rgb = RED

# ---------------------------------------------------------------------------
# Slide 9 -- TITLE_ONLY + table: efecto avalancha
# ---------------------------------------------------------------------------
s9 = prs.slides.add_slide(layouts["TITLE_ONLY"])
set_title(s9, "Efecto avalancha en acción")
add_multiline_textbox(s9, Inches(0.4), Inches(1.1), Inches(9.2), Inches(0.6), [
    '"Sobre la rama / una hoja de otoño / cae en silencio" — con y sin un punto final.',
], size=14, italic=True)

rows, cols = 3, 3
tbl_shape = s9.shapes.add_table(rows, cols, Inches(0.4), Inches(1.75), Inches(9.2), Inches(1.7))
table = tbl_shape.table
table.columns[0].width = Inches(2.4)
table.columns[1].width = Inches(3.4)
table.columns[2].width = Inches(3.4)
headers = ["Algoritmo", "Hash (texto original)", "Hash (+ un punto final)"]
data = [
    ("MD5", "c5f12d2a6fe419e912c9e57296792872", "27ae6b2b4c1808e8df03f92a06b31a43"),
    ("SHA-1", "13c3b99562aa840a4b95aa080e8add1fcfb02609", "2e4b10306f75266216133fbfd0cd503021ac6c42"),
]
for c, h in enumerate(headers):
    cell = table.cell(0, c)
    cell.text = h
    cell.text_frame.paragraphs[0].runs[0].font.bold = True
    cell.text_frame.paragraphs[0].runs[0].font.size = Pt(13)
    cell.text_frame.paragraphs[0].runs[0].font.name = FONT
    cell.fill.solid()
    cell.fill.fore_color.rgb = RED
    cell.text_frame.paragraphs[0].runs[0].font.color.rgb = WHITE
for r, row in enumerate(data, start=1):
    for c, val in enumerate(row):
        cell = table.cell(r, c)
        cell.text = val
        run = cell.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(11 if c > 0 else 13)
        run.font.name = "Courier New" if c > 0 else FONT
        run.font.color.rgb = DARK
        cell.fill.solid()
        cell.fill.fore_color.rgb = WHITE
add_multiline_textbox(s9, Inches(0.4), Inches(3.65), Inches(9.2), Inches(0.9), [
    "MD5 y SHA-1 son solo ilustrativos: ambos están rotos para uso criptográfico.",
    "DNSSEC usa SHA-256 o superior.",
], size=13, color=GRAY)

# ---------------------------------------------------------------------------
# Slide 10 -- TITLE_ONLY: Firma digital, Alice y Bob (process diagram)
# ---------------------------------------------------------------------------
s10 = prs.slides.add_slide(layouts["TITLE_ONLY"])
set_title(s10, "Firma digital: el ejemplo de Alice y Bob")
steps = [
    "1. Alice calcula el\nhash de su carta",
    "2. Alice cifra el hash\ncon su clave privada\n→ esa es la firma",
    "3. Bob recalcula el hash\ny descifra la firma con\nla clave pública de Alice",
    "4. ¿Coinciden?\nAuténtica / Alterada",
]
bw = Inches(2.05)
bh = Inches(1.5)
gap = Inches(0.28)
x0 = Inches(0.4)
by = Inches(1.9)
for i, s in enumerate(steps):
    x = x0 + i * (bw + gap)
    fill = RED if i < 3 else DARK
    add_box(s10, x, by, bw, bh, s, fill=fill, size=12)
    if i < 3:
        add_arrow(s10, x + bw, by + bh // 2, x + bw + gap, by + bh // 2)
add_multiline_textbox(s10, Inches(0.4), Inches(3.75), Inches(9.2), Inches(1.0), [
    "Esto es exactamente lo que DNSSEC aplica a los RRsets de una zona,",
    "empaquetado en un registro RRSIG.",
], size=14, color=GRAY)

# ---------------------------------------------------------------------------
# Slide 11 -- TITLE_AND_BODY: ¿Qué es DNSSEC?
# ---------------------------------------------------------------------------
s11 = prs.slides.add_slide(layouts["TITLE_AND_BODY"])
set_title(s11, "¿Qué es DNSSEC?")
body11 = get_placeholder(s11, idx=1)
set_body_bullets(body11, [
    "Extensiones al protocolo DNS que aplican firma digital a las respuestas",
    "No cifra nada — el tráfico DNS sigue siendo texto claro",
    "Permite a un resolver validar autenticidad: la respuesta viene del dueño legítimo de la zona",
    "Permite validar integridad: la respuesta no fue modificada en tránsito",
], size=18, space_after=14)

# ---------------------------------------------------------------------------
# Slide 12 -- TITLE_ONLY: cache poisoning (diagram)
# ---------------------------------------------------------------------------
s12 = prs.slides.add_slide(layouts["TITLE_ONLY"])
set_title(s12, "Qué problemas resuelve: cache poisoning")
add_multiline_textbox(s12, Inches(0.4), Inches(1.1), Inches(9.2), Inches(0.5), [
    "Ataque Kaminsky (2008): una respuesta falsa que llega antes que la legítima envenena la caché.",
], size=14)
by2 = Inches(1.85)
labels = ["Usuario", "Resolver", "Atacante\n(respuesta falsa,\nllega primero)", "IP del atacante"]
fills = [DARK, DARK, RED, RED]
for i, (lab, fill) in enumerate(zip(labels, fills)):
    x = x0 + i * (bw + gap)
    add_box(s12, x, by2, bw, bh, lab, fill=fill, size=12)
    if i < 3:
        add_arrow(s12, x + bw, by2 + bh // 2, x + bw + gap, by2 + bh // 2)
add_multiline_textbox(s12, Inches(0.4), Inches(3.75), Inches(9.2), Inches(1.0), [
    "Con DNSSEC: la respuesta falsa no está firmada con la clave privada de example.com",
    "→ la validación falla y el resolver la descarta.",
], size=14, color=RED, bold=True)

# ---------------------------------------------------------------------------
# Slide 13 -- TITLE_AND_TWO_COLUMNS: resuelve / no resuelve
# ---------------------------------------------------------------------------
s13 = prs.slides.add_slide(layouts["TITLE_AND_TWO_COLUMNS"])
set_title(s13, "Qué resuelve DNSSEC — y qué no")
colA = get_placeholder(s13, idx=1)
colB = get_placeholder(s13, idx=2)
set_body_bullets(colA, [
    "Sí resuelve",
    ("Autenticidad del origen", 1),
    ("Integridad de las respuestas", 1),
    ("Spoofing / cache poisoning", 1),
    ("Ataques on-path sobre resoluciones", 1),
], size=15, space_after=10)
colA.text_frame.paragraphs[0].runs[0].font.bold = True
colA.text_frame.paragraphs[0].runs[0].font.color.rgb = RED
set_body_bullets(colB, [
    "NO resuelve",
    ("Confidencialidad (usar DoT/DoH/DoQ)", 1),
    ("Ataques DDoS", 1),
    ("Disponibilidad (firma expirada = respuesta inválida)", 1),
    ("Contenido incorrecto publicado por el propio dueño", 1),
], size=15, space_after=10)
colB.text_frame.paragraphs[0].runs[0].font.bold = True
colB.text_frame.paragraphs[0].runs[0].font.color.rgb = GRAY

# ---------------------------------------------------------------------------
# Slide 14 -- TITLE_ONLY: Cadena de confianza (vertical chain)
# ---------------------------------------------------------------------------
s14 = prs.slides.add_slide(layouts["TITLE_ONLY"])
set_title(s14, "Cadena de confianza")
chain = [
    "Trust anchor: clave de la raíz (.)",
    "DS de .com en la raíz",
    "DNSKEY de .com",
    "DS de example.com en .com",
    "DNSKEY de example.com",
    "RRSIG sobre los registros de example.com",
]
cy = Inches(1.05)
ch = Inches(0.56)
cgap = Inches(0.09)
cw = Inches(6.0)
cx = Inches(2.0)
for i, label in enumerate(chain):
    y = cy + i * (ch + cgap)
    add_box(s14, cx, y, cw, ch, label, fill=RED if i == 0 else DARK, size=12,
            shape_type=MSO_SHAPE.ROUNDED_RECTANGLE)
    if i > 0:
        add_arrow(s14, cx + cw // 2, y - cgap, cx + cw // 2, y)
add_multiline_textbox(s14, Inches(0.4), Inches(5.05), Inches(9.2), Inches(0.5), [
    "Si un eslabón se rompe (falta un DS, una firma expiró), la validación falla para toda la zona.",
], size=12, color=GRAY, align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------------------
# Slide 15 -- TITLE_ONLY + table: Conceptos clave
# ---------------------------------------------------------------------------
s15 = prs.slides.add_slide(layouts["TITLE_ONLY"])
set_title(s15, "Conceptos clave")
tbl15_shape = s15.shapes.add_table(4, 2, Inches(0.6), Inches(1.5), Inches(8.8), Inches(2.3))
tbl15 = tbl15_shape.table
tbl15.columns[0].width = Inches(2.2)
tbl15.columns[1].width = Inches(6.6)
rows15 = [
    ("Término", "Significado breve"),
    ("Hash", "Resumen (digest) de tamaño fijo de un dato; cambia por completo si el dato cambia"),
    ("Validador", "Resolver que verifica firmas antes de aceptar una respuesta"),
    ("RRset", "Conjunto de registros del mismo nombre y tipo (lo que se firma como unidad)"),
]
for r, (a, b) in enumerate(rows15):
    for c, val in enumerate((a, b)):
        cell = tbl15.cell(r, c)
        cell.text = val
        run = cell.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(13 if r == 0 else 14)
        run.font.bold = (r == 0) or (c == 0)
        run.font.name = FONT
        run.font.color.rgb = WHITE if r == 0 else DARK
        cell.fill.solid()
        cell.fill.fore_color.rgb = RED if r == 0 else (LIGHTGRAY if c == 0 else WHITE)
add_multiline_textbox(s15, Inches(0.6), Inches(4.1), Inches(8.8), Inches(0.6), [
    "KSK, ZSK y los registros específicos de DNSSEC se cubren en detalle en el módulo 2.",
], size=13, color=GRAY, italic=True)

# ---------------------------------------------------------------------------
# Slide 16 -- TITLE_AND_BODY: Resumen
# ---------------------------------------------------------------------------
s16 = prs.slides.add_slide(layouts["TITLE_AND_BODY"])
set_title(s16, "Resumen")
body16 = get_placeholder(s16, idx=1)
set_body_bullets(body16, [
    "La firma digital con criptografía de clave pública es la base técnica de DNSSEC",
    "DNSSEC agrega autenticidad e integridad — no confidencialidad ni disponibilidad",
    "Resuelve ataques de suplantación / envenenamiento de caché con firmas verificables",
    "Depende de una cadena de confianza ininterrumpida desde la raíz hasta la zona",
    "Siguiente: Módulo 2 — los registros concretos que hacen esto posible",
], size=16, space_after=12)

# ---------------------------------------------------------------------------
# Move the closing "¡Gracias!" slide (originally slide index 2) to the end
# ---------------------------------------------------------------------------
xml_slides = prs.slides._sldIdLst
slides_list = list(xml_slides)
xml_slides.remove(slides_list[2])
xml_slides.append(slides_list[2])

prs.save("01-Introduccion-wip.pptx")
print("All slides done. Slide count:", len(prs.slides))
