#!/usr/bin/env python3
"""Build 05-Monitoreo-Troubleshooting.pptx from the lacnic46.pptx template."""
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from deck_helpers import (
    load_deck, get_placeholder, set_title, set_body_bullets, add_textbox,
    add_multiline_textbox, add_box, add_arrow, add_code_block, add_table,
    move_slide_to_end, RED, DARK, GRAY, LIGHTGRAY, WHITE, FONT, MONO,
)

prs, layouts = load_deck("lacnic46.pptx")

# ---------------------------------------------------------------------------
# Slide 1 (existing) -- Title
# ---------------------------------------------------------------------------
s1 = prs.slides[0]
s1.shapes.title.text_frame.paragraphs[0].runs[0].text = "Módulo 5: Monitoreo y troubleshooting"
sub_ph = get_placeholder(s1, idx=1)
sub_ph.text_frame.paragraphs[0].text = ""
run = sub_ph.text_frame.paragraphs[0].add_run()
run.text = "Curso de DNS y DNSSEC para Operadores  —  Carlos Martinez"
run.font.size = Pt(18)
run.font.name = FONT

# ---------------------------------------------------------------------------
# Slide 2 (existing) -- Agenda
# ---------------------------------------------------------------------------
s2 = prs.slides[1]
set_title(s2, "Agenda del módulo")
body2 = get_placeholder(s2, idx=1)
set_body_bullets(body2, [
    "Qué monitorear, y cada cuánto",
    "Árbol de diagnóstico rápido y herramientas",
    "El hueco entre el lab y un dominio real",
    "Escenario 1: firma vencida",
    "Escenario 2: DS que no coincide con la DNSKEY",
    "Escenario 3: secundario desincronizado",
    "Tabla de referencia y resumen",
], size=16, space_after=11)

# ---------------------------------------------------------------------------
# Slide 3 -- MAIN_POINT: la mayoría son incidentes operativos
# ---------------------------------------------------------------------------
s3 = prs.slides.add_slide(layouts["MAIN_POINT"])
set_title(s3, "La mayoría de los incidentes de DNSSEC\nno son ataques. Son firmas que vencieron.")

# ---------------------------------------------------------------------------
# Slide 4 -- TITLE_ONLY + table: qué monitorear
# ---------------------------------------------------------------------------
s4 = prs.slides.add_slide(layouts["TITLE_ONLY"])
set_title(s4, "Qué monitorear, y cada cuánto")
add_table(s4, Inches(0.4), Inches(1.1), Inches(9.2), Inches(4.0), [
    ("Parámetro", "Cómo chequearlo", "Frecuencia"),
    ("Vencimiento del RRSIG más próximo", "dig +dnssec, o un script", "Diario"),
    ("DS del padre vs. DNSKEY vigente", "dig DS + dnssec-dsfromkey", "Semanal / tras rollover de KSK"),
    ("Serial primario vs. secundarios", "dig SOA a cada servidor", "Cada pocos minutos"),
    ("NS y glue del padre vs. la zona", "dig NS al padre y a la zona", "Mensual / tras cambio de NS"),
    ("Rollover en curso", "rndc signing -list", "Según la dnssec-policy"),
], col_widths=[Inches(3.1), Inches(3.3), Inches(2.8)], size=11, header_size=13,
    first_col_bold=True)

# ---------------------------------------------------------------------------
# Slide 5 -- TITLE_ONLY: árbol de diagnóstico rápido
# ---------------------------------------------------------------------------
s5 = prs.slides.add_slide(layouts["TITLE_ONLY"])
set_title(s5, "Árbol de diagnóstico rápido")
add_multiline_textbox(s5, Inches(0.4), Inches(1.05), Inches(9.2), Inches(0.4), [
    "Reporte: \"example.com no resuelve\" / SERVFAIL",
], size=14, bold=True, color=RED)
# Row 1: dig sin +dnssec
by1 = Inches(1.55)
bh = Inches(0.62)
add_box(s5, Inches(0.4), by1, Inches(9.2), bh, "dig sin +dnssec — ¿responde?", fill=DARK, size=13)
by2 = Inches(2.35)
add_box(s5, Inches(0.4), by2, Inches(4.4), bh, "No responde nada\n→ servicio caído/inalcanzable (módulo 3)", fill=GRAY, size=11)
add_box(s5, Inches(5.2), by2, Inches(4.4), bh, "Responde datos\n→ correr delv @servidor A", fill=RED, size=11)
by3 = Inches(3.15)
labels3 = [
    "\"RRSIG has expired\"\n→ Escenario 1",
    "DS no coincide\n→ Escenario 2",
    "\"insecure\"\n→ falta DS (módulo 4)",
]
bw3 = Inches(2.95)
gap3 = Inches(0.18)
for i, lab in enumerate(labels3):
    x = Inches(0.4) + i * (bw3 + gap3)
    add_box(s5, x, by3, bw3, bh, lab, fill=DARK, size=11)
by4 = Inches(3.95)
add_box(s5, Inches(0.4), by4, Inches(9.2), bh, "Serial del secundario distinto al primario → transferencia rota — Escenario 3", fill=RED, size=12)

# ---------------------------------------------------------------------------
# Slide 6 -- TITLE_ONLY + table: herramientas
# ---------------------------------------------------------------------------
s6 = prs.slides.add_slide(layouts["TITLE_ONLY"])
set_title(s6, "Herramientas")
add_table(s6, Inches(0.4), Inches(1.2), Inches(9.2), Inches(2.6), [
    ("Herramienta", "Para qué"),
    ("dig +dnssec", "Ver RRSIG/DNSKEY tal cual los recibe un cliente"),
    ("delv", "Validar la cadena de confianza completa, como un resolver validador"),
    ("rndc zonestatus / signing -list", "Estado de firmado sin ir al log"),
    ("named-checkzone / dnssec-verify", "Validar sintaxis y firma antes de confiar en que \"debería andar\""),
], col_widths=[Inches(3.2), Inches(6.0)], size=13, header_size=14, first_col_bold=True,
    mono_cols={0})

# ---------------------------------------------------------------------------
# Slide 7 -- TITLE_ONLY: script de alerta
# ---------------------------------------------------------------------------
s7 = prs.slides.add_slide(layouts["TITLE_ONLY"])
set_title(s7, "Un script simple de alerta")
add_code_block(s7, Inches(0.4), Inches(1.1), Inches(9.2), Inches(3.5), [
    "#!/bin/bash",
    "# chequea-rrsig.sh — alerta si una firma vence en menos de N días",
    "ZONA=\"example.com\"; UMBRAL_DIAS=3",
    "",
    "VENCE=$(dig +dnssec @127.0.0.1 \"$ZONA\" A | awk '/RRSIG/ {print $9; exit}')",
    "DIAS_RESTANTES=$(( (VENCE_EPOCH - AHORA_EPOCH) / 86400 ))",
    "",
    "if [ \"$DIAS_RESTANTES\" -lt \"$UMBRAL_DIAS\" ]; then",
    "    echo \"ALERTA: RRSIG de $ZONA vence en $DIAS_RESTANTES día(s)\" | logger",
    "fi",
], size=13)
add_multiline_textbox(s7, Inches(0.4), Inches(4.75), Inches(9.2), Inches(0.7), [
    "Corrido por cron una vez al día — cubre el escenario más común de todos.",
], size=13, color=GRAY, italic=True)

# ---------------------------------------------------------------------------
# Slide 8 -- TITLE_AND_BODY: herramientas externas
# ---------------------------------------------------------------------------
s8 = prs.slides.add_slide(layouts["TITLE_AND_BODY"])
set_title(s8, "Herramientas externas")
body8 = get_placeholder(s8, idx=1)
set_body_bullets(body8, [
    "DNSViz — grafica la cadena de confianza completa, marca en rojo el eslabón que falla",
    "Zonemaster — batería de chequeos más amplia: delegación, glue, consistencia entre servidores",
    "Monitoreo continuo — statistics-channels + bind_exporter (Prometheus), templates de Zabbix",
], size=17, space_after=14)

# ---------------------------------------------------------------------------
# Slide 9 -- TITLE_ONLY: el hueco entre el lab y un dominio real
# ---------------------------------------------------------------------------
s9 = prs.slides.add_slide(layouts["TITLE_ONLY"])
set_title(s9, "El hueco entre el lab y un dominio real")
add_code_block(s9, Inches(0.4), Inches(1.15), Inches(9.2), Inches(1.3), [
    "$ dig @1.1.1.1 example.com A         # como lo ve un resolver validador",
    "$ dig @1.1.1.1 example.com A +cd     # validación desactivada, para comparar",
], size=13)
add_multiline_textbox(s9, Inches(0.4), Inches(2.75), Inches(9.2), Inches(1.2), [
    "Si la primera falla (o da SERVFAIL) y la segunda funciona, el problema es",
    "específicamente de DNSSEC — no de DNS en general.",
], size=15)

# ---------------------------------------------------------------------------
# Slide 10 -- SECTION_TITLE_AND_DESCRIPTION: Tres escenarios
# ---------------------------------------------------------------------------
s10 = prs.slides.add_slide(layouts["SECTION_TITLE_AND_DESCRIPTION"])
set_title(s10, "Tres escenarios")
sub10 = get_placeholder(s10, idx=1)
r = sub10.text_frame.paragraphs[0].add_run()
r.text = "Reproducidos deliberadamente sobre el mismo example.com"
r.font.size = Pt(18)
r.font.name = FONT
body10 = get_placeholder(s10, idx=2)
set_body_bullets(body10, ["Firma vencida", "DS que no coincide con la DNSKEY", "Secundario desincronizado"], size=17)

# ---------------------------------------------------------------------------
# Slide 11 -- TITLE_AND_BODY: Escenario 1 contexto
# ---------------------------------------------------------------------------
s11 = prs.slides.add_slide(layouts["TITLE_AND_BODY"])
set_title(s11, "Escenario 1: firma vencida")
body11 = get_placeholder(s11, idx=1)
set_body_bullets(body11, [
    "La causa más común en la práctica: named estuvo caído más tiempo del que dura la validez de las firmas",
    "Una firma tiene fecha de vencimiento absoluta — no le importa si el servidor estuvo vivo",
    "Para reproducirlo en minutos, achicamos la política solo para este ejercicio",
], size=17, space_after=13)

# ---------------------------------------------------------------------------
# Slide 12 -- TITLE_ONLY: reproducir el escenario 1
# ---------------------------------------------------------------------------
s12 = prs.slides.add_slide(layouts["TITLE_ONLY"])
set_title(s12, "Reproducir la caída")
add_code_block(s12, Inches(0.4), Inches(1.1), Inches(9.2), Inches(2.6), [
    "signatures-validity PT10M;   // normalmente P2W",
    "signatures-refresh PT5M;     // normalmente P5D",
    "",
    "$ sudo rndc reconfig",
    "$ sudo systemctl stop bind9",
    "$ sleep 720                  # 12 minutos — más que la validez",
    "$ sudo systemctl start bind9",
], size=13)
add_multiline_textbox(s12, Inches(0.4), Inches(4.0), Inches(9.2), Inches(0.9), [
    "named vuelve a levantar sirviendo la última zona firmada que tenía en disco",
    "— con la firma ya vencida, hasta que su ciclo de mantenimiento la renueve.",
], size=13, color=GRAY)

# ---------------------------------------------------------------------------
# Slide 13 -- TITLE_ONLY: delv SERVFAIL y recuperación
# ---------------------------------------------------------------------------
s13 = prs.slides.add_slide(layouts["TITLE_ONLY"])
set_title(s13, "SERVFAIL — y la recuperación automática")
add_code_block(s13, Inches(0.4), Inches(1.1), Inches(9.2), Inches(1.7), [
    "$ delv @127.0.0.1 example.com A",
    ";; verify failed due to bad signature: RRSIG has expired",
    ";; broken trust chain resolving 'example.com/A/IN'",
    ";; resolution failed: SERVFAIL",
], size=13)
add_code_block(s13, Inches(0.4), Inches(3.05), Inches(9.2), Inches(1.15), [
    "$ delv @127.0.0.1 example.com A     # named renovó la firma solo",
    "; fully validated",
    "93.184.216.34",
], size=13)
add_multiline_textbox(s13, Inches(0.4), Inches(4.4), Inches(9.2), Inches(0.6), [
    "No olvidar revertir la política a los valores reales del módulo 4 después del ejercicio.",
], size=12, color=RED, italic=True)

# ---------------------------------------------------------------------------
# Slide 14 -- TITLE_AND_BODY: Escenario 2 contexto
# ---------------------------------------------------------------------------
s14 = prs.slides.add_slide(layouts["TITLE_AND_BODY"])
set_title(s14, "Escenario 2: DS que no coincide con la DNSKEY")
body14 = get_placeholder(s14, idx=1)
set_body_bullets(body14, [
    "Reproduce el error más común al publicar el DS: un carácter mal copiado al registrador",
    "delv -a permite validar contra un trust anchor declarado a mano, sin depender del padre real",
    "Mismo síntoma que un rollover de KSK donde el DS no se actualizó a tiempo",
], size=17, space_after=13)

# ---------------------------------------------------------------------------
# Slide 15 -- TITLE_AND_TWO_COLUMNS: DS correcto vs incorrecto
# ---------------------------------------------------------------------------
s15 = prs.slides.add_slide(layouts["TITLE_AND_TWO_COLUMNS"])
set_title(s15, "El mismo delv, dos trust anchors")
col1 = get_placeholder(s15, idx=1)
col2 = get_placeholder(s15, idx=2)
tf1 = col1.text_frame
tf1.word_wrap = True
p = tf1.paragraphs[0]
pPr = p._pPr
if pPr is not None:
    p._p.remove(pPr)
r = p.add_run(); r.text = "DS correcto"; r.font.bold = True; r.font.size = Pt(16); r.font.name = FONT; r.font.color.rgb = RED
add_code_block(s15, Inches(0.5), Inches(2.0), Inches(4.2), Inches(1.4), [
    "$ delv -a ta-correcto.conf ...",
    "; fully validated",
    "93.184.216.34",
], size=11)
tf2 = col2.text_frame
tf2.word_wrap = True
p2 = tf2.paragraphs[0]
pPr2 = p2._pPr
if pPr2 is not None:
    p2._p.remove(pPr2)
r2 = p2.add_run(); r2.text = "DS con un dígito alterado"; r2.font.bold = True; r2.font.size = Pt(16); r2.font.name = FONT; r2.font.color.rgb = RED
add_code_block(s15, Inches(5.1), Inches(2.0), Inches(4.4), Inches(1.4), [
    "$ delv -a ta-incorrecto.conf ...",
    ";; verify failed: DS does not",
    "   match DNSKEY",
    ";; resolution failed: SERVFAIL",
], size=11)
add_multiline_textbox(s15, Inches(0.5), Inches(3.7), Inches(9.0), Inches(1.2), [
    "dnssec-verify sobre la zona sigue diciendo que todo firma perfecto — el",
    "problema no está en la zona, está en lo que el padre publica.",
], size=14, color=GRAY)

# ---------------------------------------------------------------------------
# Slide 16 -- TITLE_ONLY: Escenario 3, secundario desincronizado
# ---------------------------------------------------------------------------
s16 = prs.slides.add_slide(layouts["TITLE_ONLY"])
set_title(s16, "Escenario 3: secundario desincronizado")
add_code_block(s16, Inches(0.4), Inches(1.1), Inches(9.2), Inches(2.0), [
    "$ dig @192.0.2.1 example.com SOA +short   # primario",
    "ns1.example.com. admin.example.com. 2026080502 ...",
    "",
    "$ dig @192.0.2.2 example.com SOA +short   # secundario",
    "ns1.example.com. admin.example.com. 2026080401 ...   # ¡serial viejo!",
], size=13)
add_multiline_textbox(s16, Inches(0.4), Inches(3.35), Inches(9.2), Inches(1.4), [
    "La transferencia falla en silencio — el secundario sigue respondiendo con su",
    "última copia buena, sin avisar por sí solo que quedó atrás. Combinación",
    "peligrosa con el escenario 1: firmas cada vez más viejas, sin que nadie note.",
], size=14, color=GRAY)

# ---------------------------------------------------------------------------
# Slide 17 -- TITLE_ONLY: diagnóstico y arreglo
# ---------------------------------------------------------------------------
s17 = prs.slides.add_slide(layouts["TITLE_ONLY"])
set_title(s17, "Diagnóstico y arreglo")
add_code_block(s17, Inches(0.4), Inches(1.1), Inches(9.2), Inches(2.3), [
    "$ sudo rndc retransfer example.com   # en el secundario",
    "... failed to connect: timed out",
    "",
    "# resuelto el ACL/firewall:",
    "$ sudo rndc retransfer example.com",
    "$ dig @192.0.2.2 example.com SOA +short",
    "ns1.example.com. admin.example.com. 2026080502 ...   # sincronizado",
], size=13)
add_multiline_textbox(s17, Inches(0.4), Inches(3.7), Inches(9.2), Inches(0.9), [
    "Mismo orden de diagnóstico que en el módulo 3: ACL/TSIG del primario,",
    "conectividad, logs de ambos lados.",
], size=14, color=GRAY)

# ---------------------------------------------------------------------------
# Slide 18 -- TITLE_ONLY + table: otros síntomas comunes
# ---------------------------------------------------------------------------
s18 = prs.slides.add_slide(layouts["TITLE_ONLY"])
set_title(s18, "Tabla de referencia: otros síntomas")
add_table(s18, Inches(0.3), Inches(1.0), Inches(9.4), Inches(3.9), [
    ("Síntoma", "Causa probable", "Solución"),
    ("Delegación \"lame\"", "NS del padre no coincide con la zona", "Corregir NS/glue en el registrador"),
    ("Falla en algunos resolvers", "TTL viejo del DS/DNSKEY en caché", "Esperar el TTL — es propagación"),
    ("No vuelve a firmar tras editar", "Se perdió la dnssec-policy/inline-signing", "Revisar rndc zonestatus"),
    ("rndc no responde tras cambio de hora", "Reloj desincronizado", "Sincronizar con NTP"),
    ("Algoritmo nuevo no valida en algunos resolvers", "Resolver viejo, no lo soporta aún", "Evaluar esperar antes de migrar"),
], col_widths=[Inches(3.0), Inches(3.3), Inches(3.1)], size=10.5, header_size=12,
    first_col_bold=True)

# ---------------------------------------------------------------------------
# Slide 19 -- TITLE_AND_BODY: Resumen
# ---------------------------------------------------------------------------
s19 = prs.slides.add_slide(layouts["TITLE_AND_BODY"])
set_title(s19, "Resumen")
body19 = get_placeholder(s19, idx=1)
set_body_bullets(body19, [
    "delv distingue insecure, SERVFAIL por firma vencida, y SERVFAIL por DS que no coincide",
    "Una zona puede firmar perfecto (dnssec-verify OK) y aun así fallar la validación para todos",
    "Un secundario desincronizado no avisa solo — hay que vigilar el serial activamente",
    "Siguiente: Módulo 6 — separar el firmado de la publicación en servidores distintos",
], size=16, space_after=13)

# ---------------------------------------------------------------------------
# Move the closing "¡Gracias!" slide to the end
# ---------------------------------------------------------------------------
move_slide_to_end(prs, 2)

prs.save("05-Monitoreo-Troubleshooting-wip.pptx")
print("Module 5 done. Slide count:", len(prs.slides))
