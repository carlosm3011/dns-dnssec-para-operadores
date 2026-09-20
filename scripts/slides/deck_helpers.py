#!/usr/bin/env python3
"""Shared helpers for building LACNIC-template (lacnic46.pptx) course decks."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

RED = RGBColor(0xC5, 0x34, 0x25)
DARK = RGBColor(0x33, 0x33, 0x33)
GRAY = RGBColor(0x66, 0x66, 0x66)
LIGHTGRAY = RGBColor(0xF3, 0xF3, 0xF3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"
MONO = "Courier New"


def load_deck(template_path="lacnic46.pptx"):
    prs = Presentation(template_path)
    layouts = {l.name: l for l in prs.slide_masters[0].slide_layouts}
    return prs, layouts


def get_placeholder(slide, ph_type=None, idx=None):
    for ph in slide.placeholders:
        if idx is not None and ph.placeholder_format.idx == idx:
            return ph
        if ph_type is not None and str(ph.placeholder_format.type) == ph_type:
            return ph
    return None


def set_title(slide, text):
    slide.shapes.title.text_frame.text = text


def set_body_bullets(placeholder, items, size=16, space_after=10, bullet_first=True):
    """items: list of str, or (str, level) tuples.

    bullet_first: when True (flat bullet lists), strip any inherited
    <a:buNone/> on paragraph 0 so the first line bullets like the rest.
    Set False when item 0 is meant to read as an un-bulleted header
    (e.g. two-column layouts where the first line is a column title).
    """
    tf = placeholder.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if i == 0 and bullet_first:
            # Some pre-existing template placeholders carry an explicit
            # <a:pPr marL="0" indent="0" ...><a:buNone/></a:pPr> baked onto
            # their first (only) paragraph, which both kills the bullet and
            # zeroes the indent that gives it its hanging space. Drop the
            # whole override so this paragraph is built fresh, matching the
            # ones we add below via p.level.
            pPr = p._pPr
            if pPr is not None:
                p._p.remove(pPr)
        p.level = level
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size - (2 if level > 0 else 0))
        run.font.name = FONT
        run.font.color.rgb = DARK
        p.space_after = Pt(space_after)


def add_textbox(slide, left, top, width, height, text, size=14, bold=False,
                 color=DARK, align=PP_ALIGN.LEFT, font=FONT, italic=False):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font
    run.font.color.rgb = color
    return tb


def add_multiline_textbox(slide, left, top, width, height, lines, size=13,
                           color=DARK, font=FONT, line_spacing=1.15, align=PP_ALIGN.LEFT,
                           italic=False, bold=False):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.name = font
        run.font.color.rgb = color
        run.font.italic = italic
        run.font.bold = bold
        p.line_spacing = line_spacing
    return tb


def add_box(slide, left, top, width, height, text, fill=RED, text_color=WHITE,
            size=13, bold=True, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE):
    shp = slide.shapes.add_shape(shape_type, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(45700)
    tf.margin_right = Emu(45700)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = FONT
    run.font.color.rgb = text_color
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return shp


def add_arrow(slide, x1, y1, x2, y2, color=GRAY, width=Pt(2)):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    conn.line.color.rgb = color
    conn.line.width = width
    ln = conn.line._get_or_add_ln()
    tail = ln.makeelement(qn('a:tailEnd'), {'type': 'triangle'})
    ln.append(tail)
    return conn


def add_code_block(slide, left, top, width, height, lines, size=12, fill=LIGHTGRAY,
                    color=DARK, prompt_color=RED):
    """Rounded rect with monospaced lines. Lines starting with '$ ' get the
    prompt colored to stand out (shell commands vs. output)."""
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    box.fill.solid()
    box.fill.fore_color.rgb = fill
    box.line.fill.background()
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_top = Inches(0.18)
    tf.margin_bottom = Inches(0.18)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        if line.startswith("$ "):
            r1 = p.add_run()
            r1.text = "$ "
            r1.font.size = Pt(size)
            r1.font.name = MONO
            r1.font.bold = True
            r1.font.color.rgb = prompt_color
            r2 = p.add_run()
            r2.text = line[2:]
            r2.font.size = Pt(size)
            r2.font.name = MONO
            r2.font.color.rgb = color
        else:
            r = p.add_run()
            r.text = line
            r.font.size = Pt(size)
            r.font.name = MONO
            r.font.color.rgb = color
    return box


def add_table(slide, left, top, width, height, rows, header=True, col_widths=None,
              header_fill=RED, header_color=WHITE, body_color=DARK, size=13,
              header_size=13, first_col_bold=False, mono_cols=None):
    """rows: list of tuples of strings. First row is header if header=True."""
    nrows = len(rows)
    ncols = len(rows[0])
    shape = slide.shapes.add_table(nrows, ncols, left, top, width, height)
    table = shape.table
    if col_widths:
        for c, w in enumerate(col_widths):
            table.columns[c].width = w
    mono_cols = mono_cols or set()
    for r, row in enumerate(rows):
        is_header = header and r == 0
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(val)
            run = cell.text_frame.paragraphs[0].runs[0]
            run.font.size = Pt(header_size if is_header else size)
            run.font.name = MONO if c in mono_cols and not is_header else FONT
            run.font.bold = is_header or (first_col_bold and c == 0)
            run.font.color.rgb = header_color if is_header else body_color
            cell.fill.solid()
            if is_header:
                cell.fill.fore_color.rgb = header_fill
            else:
                cell.fill.fore_color.rgb = LIGHTGRAY if c == 0 and first_col_bold else WHITE
    return table


def move_slide_to_end(prs, index):
    xml_slides = prs.slides._sldIdLst
    slides_list = list(xml_slides)
    xml_slides.remove(slides_list[index])
    xml_slides.append(slides_list[index])
