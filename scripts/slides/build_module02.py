#!/usr/bin/env python3
"""Build 02-Registros.pptx from the lacnic46.pptx template."""
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from deck_helpers import (
    load_deck, get_placeholder, set_title, set_body_bullets, add_textbox,
    add_multiline_textbox, add_box, add_arrow, add_code_block, add_table,
    move_slide_to_end, RED, DARK, GRAY, LIGHTGRAY, WHITE, FONT, MONO,
)

prs, layouts = load_deck("lacnic46.pptx")

# ---------------------------------------------------------------------------
# Slide 1 (existing) -- Title
# ---------------------------------------------------------------------------
s1 = prs.slides[0]
s1.shapes.title.text_frame.paragraphs[0].runs[0].text = "Módulo 2: Los registros de DNSSEC"
sub_ph = get_placeholder(s1, idx=1)
sub_ph.text_frame.paragraphs[0].text = ""
run = sub_ph.text_frame.paragraphs[0].add_run()
run.text = "Curso de DNS y DNSSEC para Operadores  —  Carlos Martinez"
run.font.size = Pt(18)
run.font.name = FONT

# ---------------------------------------------------------------------------
# Slide 2 (existing) -- Agenda
# ---------------------------------------------------------------------------
s2 = prs.slides[1]
set_title(s2, "Agenda del módulo")
body2 = get_placeholder(s2, idx=1)
set_body_bullets(body2, [
    "DNSKEY — la clave pública de la zona",
    "RRSIG — la firma sobre un RRset",
    "DS — el enlace hacia el padre",
    "NSEC — negar que algo no existe, de forma autenticada",
    "Resumen",
], size=19, space_after=16)

# ---------------------------------------------------------------------------
# Slide 3 -- SECTION_TITLE_AND_DESCRIPTION: Cuatro registros nuevos
# ---------------------------------------------------------------------------
s3 = prs.slides.add_slide(layouts["SECTION_TITLE_AND_DESCRIPTION"])
set_title(s3, "Cuatro registros nuevos")
sub3 = get_placeholder(s3, idx=1)
r = sub3.text_frame.paragraphs[0].add_run()
r.text = "Cada uno cumple un rol distinto en la cadena de confianza"
r.font.size = Pt(18)
r.font.name = FONT
body3 = get_placeholder(s3, idx=2)
set_body_bullets(body3, ["DNSKEY", "RRSIG", "DS", "NSEC"], size=17)

# ---------------------------------------------------------------------------
# Slide 4 -- TITLE_AND_BODY: DNSKEY
# ---------------------------------------------------------------------------
s4 = prs.slides.add_slide(layouts["TITLE_AND_BODY"])
set_title(s4, "DNSKEY — la clave pública de la zona")
body4 = get_placeholder(s4, idx=1)
set_body_bullets(body4, [
    "Contiene la clave pública de la zona",
    "Se publica en la propia zona — cualquier validador puede obtenerla",
    "isc.org publica dos DNSKEY: una ZSK (Zone Signing Key) y una KSK (Key Signing Key)",
    "El motivo de usar dos claves en vez de una se ve en el módulo 4",
], size=18, space_after=14)

# ---------------------------------------------------------------------------
# Slide 5 -- TITLE_ONLY: ejemplo dig DNSKEY
# ---------------------------------------------------------------------------
s5 = prs.slides.add_slide(layouts["TITLE_ONLY"])
set_title(s5, "dig DNSKEY isc.org")
add_code_block(s5, Inches(0.4), Inches(1.15), Inches(9.2), Inches(3.4), [
    "$ dig +dnssec +multiline DNSKEY isc.org",
    "",
    "isc.org.  405 IN DNSKEY 256 3 13 ( ... )",
    "          ; ZSK; alg = ECDSAP256SHA256 ; key id = 27566",
    "isc.org.  405 IN DNSKEY 257 3 13 ( ... )",
    "          ; KSK; alg = ECDSAP256SHA256 ; key id = 7250",
], size=14)
add_multiline_textbox(s5, Inches(0.4), Inches(4.75), Inches(9.2), Inches(0.7), [
    "dig anota el rol (ZSK/KSK) y el key id de cada clave entre paréntesis.",
], size=13, color=GRAY, italic=True)

# ---------------------------------------------------------------------------
# Slide 6 -- TITLE_AND_BODY: RRSIG
# ---------------------------------------------------------------------------
s6 = prs.slides.add_slide(layouts["TITLE_AND_BODY"])
set_title(s6, "RRSIG — la firma sobre un RRset")
body6 = get_placeholder(s6, idx=1)
set_body_bullets(body6, [
    "La firma digital del módulo 1, empaquetada como registro DNS",
    "Cada RRset firmado tiene su propio RRSIG",
    "Campos clave: algoritmo, vencimiento/inicio de validez, key id de la clave firmante, nombre de la zona",
], size=18, space_after=14)

# ---------------------------------------------------------------------------
# Slide 7 -- TITLE_ONLY: ejemplo dig A + RRSIG
# ---------------------------------------------------------------------------
s7 = prs.slides.add_slide(layouts["TITLE_ONLY"])
set_title(s7, "dig A isc.org — con su RRSIG")
add_code_block(s7, Inches(0.4), Inches(1.15), Inches(9.2), Inches(2.7), [
    "$ dig +dnssec +multiline A isc.org",
    "",
    "isc.org.  300 IN A 151.101.2.217",
    "isc.org.  300 IN RRSIG A 13 2 300 (",
    "    20260719192313 20260705182333 27566 isc.org.",
    "    ...)",
], size=13)
add_multiline_textbox(s7, Inches(0.4), Inches(4.05), Inches(9.2), Inches(1.4), [
    "13 = algoritmo   ·   20260719192313 / 20260705182333 = vencimiento / inicio",
    "27566 = key id de la ZSK que firmó   ·   isc.org. = zona firmante",
], size=13, color=GRAY)

# ---------------------------------------------------------------------------
# Slide 8 -- TITLE_AND_TWO_COLUMNS: KSK vs ZSK, quién firma qué
# ---------------------------------------------------------------------------
s8 = prs.slides.add_slide(layouts["TITLE_AND_TWO_COLUMNS"])
set_title(s8, "Quién firma qué")
col1 = get_placeholder(s8, idx=1)
col2 = get_placeholder(s8, idx=2)
set_body_bullets(col1, [
    "KSK (key id 7250)",
    ("Firma únicamente el DNSKEY RRset", 1),
], size=17, space_after=12, bullet_first=False)
col1.text_frame.paragraphs[0].runs[0].font.bold = True
col1.text_frame.paragraphs[0].runs[0].font.color.rgb = RED
set_body_bullets(col2, [
    "ZSK (key id 27566)",
    ("Firma todo lo demás: A, SOA, NSEC, etc.", 1),
], size=17, space_after=12, bullet_first=False)
col2.text_frame.paragraphs[0].runs[0].font.bold = True
col2.text_frame.paragraphs[0].runs[0].font.color.rgb = RED

# ---------------------------------------------------------------------------
# Slide 9 -- TITLE_AND_BODY: DS
# ---------------------------------------------------------------------------
s9 = prs.slides.add_slide(layouts["TITLE_AND_BODY"])
set_title(s9, "DS — el enlace hacia el padre")
body9 = get_placeholder(s9, idx=1)
set_body_bullets(body9, [
    "No vive en la zona misma, sino en la zona padre",
    "Es un hash de la KSK del hijo",
    "Extiende la cadena de confianza de un eslabón al siguiente",
], size=18, space_after=14)

# ---------------------------------------------------------------------------
# Slide 10 -- TITLE_ONLY: ejemplo dig DS + notas
# ---------------------------------------------------------------------------
s10 = prs.slides.add_slide(layouts["TITLE_ONLY"])
set_title(s10, "dig DS isc.org")
add_code_block(s10, Inches(0.4), Inches(1.1), Inches(9.2), Inches(1.6), [
    "$ dig +dnssec +multiline DS isc.org",
    "",
    "isc.org.  3600 IN DS 7250 13 2 ( ... )",
    "isc.org.  3600 IN RRSIG DS 8 2 3600 (",
    "    20260730154208 20260709144208 13950 org. ...)",
], size=13)
add_multiline_textbox(s10, Inches(0.4), Inches(2.9), Inches(9.2), Inches(2.2), [
    "•  7250 = key id de la KSK de isc.org — el DS apunta a esa clave específica",
    "•  Firmado por la clave 13950, que pertenece a org. — el padre certifica al hijo",
    "•  Un validador hashea la KSK recibida y la compara contra este DS",
], size=14)

# ---------------------------------------------------------------------------
# Slide 11 -- MAIN_POINT: el problema de NXDOMAIN sin firmar
# ---------------------------------------------------------------------------
s11 = prs.slides.add_slide(layouts["MAIN_POINT"])
set_title(s11, "Sin DNSSEC, un NXDOMAIN no está firmado:\nun atacante podría falsificarlo.")

# ---------------------------------------------------------------------------
# Slide 12 -- TITLE_ONLY: NSEC, ejemplo y explicación
# ---------------------------------------------------------------------------
s12 = prs.slides.add_slide(layouts["TITLE_ONLY"])
set_title(s12, "NSEC — negar de forma autenticada")
add_code_block(s12, Inches(0.4), Inches(1.1), Inches(9.2), Inches(1.5), [
    "$ dig +dnssec +multiline A doesnotexist.isc.org",
    "",
    "docs.isc.org. 3600 IN NSEC dommel.isc.org. A RRSIG NSEC",
    "docs.isc.org. 3600 IN RRSIG NSEC 13 3 3600 ( ... )",
], size=13)
add_multiline_textbox(s12, Inches(0.4), Inches(2.85), Inches(9.2), Inches(2.0), [
    "\"No existe ningún nombre entre este registro y el siguiente, en orden",
    "alfabético.\" doesnotexist.isc.org cae entre docs. y dommel. — el NSEC",
    "firmado prueba criptográficamente la ausencia, no solo un NXDOMAIN sin firmar.",
], size=14)
add_multiline_textbox(s12, Inches(0.4), Inches(4.35), Inches(9.2), Inches(0.9), [
    "NSEC3 — misma lógica, con nombres hasheados para evitar zone walking (zonas grandes o sensibles a enumeración).",
], size=13, color=GRAY, italic=True)

# ---------------------------------------------------------------------------
# Slide 13 -- TITLE_ONLY + table: Resumen
# ---------------------------------------------------------------------------
s13 = prs.slides.add_slide(layouts["TITLE_ONLY"])
set_title(s13, "Resumen: los cuatro registros")
add_table(s13, Inches(0.4), Inches(1.3), Inches(9.2), Inches(2.6), [
    ("Registro", "Vive en", "Firma / protege"),
    ("DNSKEY", "La propia zona", "Publica las claves públicas (ZSK/KSK)"),
    ("RRSIG", "La propia zona", "La firma de un RRset puntual"),
    ("DS", "La zona padre", "El hash de la KSK del hijo — extiende la cadena"),
    ("NSEC", "La propia zona", "Prueba autenticada de que un nombre/tipo no existe"),
], col_widths=[Inches(1.4), Inches(2.0), Inches(5.8)], size=13, header_size=14,
    first_col_bold=True)

# ---------------------------------------------------------------------------
# Slide 14 -- TITLE_AND_BODY: Resumen (bullets)
# ---------------------------------------------------------------------------
s14 = prs.slides.add_slide(layouts["TITLE_AND_BODY"])
set_title(s14, "Resumen")
body14 = get_placeholder(s14, idx=1)
set_body_bullets(body14, [
    "Todo RRset firmado tiene su RRSIG; el DNSKEY RRset lo firma la KSK, el resto la ZSK",
    "El DS es el único registro DNSSEC que no vive en la zona que protege",
    "NSEC y NSEC3 resuelven el mismo problema de denegación autenticada",
    "Siguiente: Módulo 3 — levantar un servidor BIND y publicar una zona autoritativa",
], size=17, space_after=14)

# ---------------------------------------------------------------------------
# Move the closing "¡Gracias!" slide to the end
# ---------------------------------------------------------------------------
move_slide_to_end(prs, 2)

prs.save("02-Registros-wip.pptx")
print("Module 2 done. Slide count:", len(prs.slides))
